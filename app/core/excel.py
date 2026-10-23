"""Excel: шаблон для імпорту, імпорт клієнтів/опонентів, експорт клієнтів і справ."""

import io
from dataclasses import dataclass, field
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core import crypto
from app.core.conflicts import normalize_name, normalize_phone
from app.core.defaults import CATEGORIES
from app.core.statuses import CaseStatus, ConflictStatus
from app.db.models import Case, Client, Party

TEMPLATE_COLUMNS = [
    "ПІБ / назва",
    "Тип (фізособа/юрособа)",
    "Телефон",
    "РНОКПП / ЄДРПОУ",
    "Роль (клієнт/опонент/інша сторона)",
    "Справа / пов'язаний клієнт",
    "Примітка",
]

EXAMPLE_ROWS = [
    ["Коваленко Марія Іванівна", "фізособа", "+380501234567", "", "клієнт", "Розлучення 2025", ""],
    ["Коваленко Петро Олегович", "фізособа", "", "", "опонент", "Коваленко Марія Іванівна", ""],
    ["ТОВ «Приклад»", "юрособа", "", "12345678", "інша сторона", "", "Приклад рядка — видаліть"],
]

HEADER_FILL = PatternFill("solid", fgColor="DDE6F0")


def _style_sheet(ws, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def _to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Клієнти"
    ws.append(TEMPLATE_COLUMNS)
    for row in EXAMPLE_ROWS:
        ws.append(row)
    _style_sheet(ws, [34, 14, 18, 16, 22, 30, 30])
    return _to_bytes(wb)


@dataclass
class ImportResult:
    clients: int = 0
    parties: int = 0
    skipped: int = 0
    errors: list[str] = field(default_factory=list)


def _cell(row: tuple, idx: int) -> str:
    if idx >= len(row) or row[idx] is None:
        return ""
    value = row[idx]
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


async def import_xlsx(session: AsyncSession, data: bytes) -> ImportResult:
    result = ImportResult()
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        result.errors.append(f"Не вдалося відкрити файл: {exc}")
        return result
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(c or "").strip().lower() for c in next(rows, ())]
    if not header or not header[0].startswith("піб"):
        result.errors.append("Перший рядок має бути заголовком шаблону (перша колонка «ПІБ / назва»).")
        return result

    existing = (await session.scalars(select(Client))).all()
    seen = {(c.name_norm, c.phone) for c in existing}
    seen_codes = {c.tax_id_hash for c in existing if c.tax_id_hash}

    for line_no, row in enumerate(rows, start=2):
        name = _cell(row, 0)
        if not name:
            continue
        kind = _cell(row, 1).lower()
        phone = normalize_phone(_cell(row, 2))
        code = "".join(ch for ch in _cell(row, 3) if ch.isdigit())
        role = _cell(row, 4).lower() or "клієнт"
        related = _cell(row, 5)
        note = _cell(row, 6)
        is_company = kind.startswith("юр") or len(code) == 8
        if code and len(code) not in (8, 10):
            result.errors.append(
                f"Рядок {line_no}: код «{code}» має бути 10 (РНОКПП) або 8 (ЄДРПОУ) цифр — пропущено код."
            )
            code = ""
        code_hash = crypto.code_hash(code) if code else ""
        norm = normalize_name(name)

        if role.startswith("клієнт"):
            if (norm, phone) in seen or (code_hash and code_hash in seen_codes):
                result.skipped += 1
                continue
            session.add(
                Client(
                    full_name=name,
                    name_norm=norm,
                    is_company=is_company,
                    phone=phone,
                    tax_id_enc=crypto.encrypt(code) if code else "",
                    tax_id_hash=code_hash,
                    source="import",
                    has_contract=True,
                    note="\n".join(x for x in (related, note) if x),
                )
            )
            seen.add((norm, phone))
            if code_hash:
                seen_codes.add(code_hash)
            result.clients += 1
        else:
            session.add(
                Party(
                    name=name,
                    name_norm=norm,
                    is_company=is_company,
                    tax_id_hash=code_hash,
                    role="опонент" if role.startswith("опон") else (role or "інша сторона"),
                    related_client_name=related,
                    note=note,
                )
            )
            result.parties += 1
    return result


def _fmt(dt: datetime | None) -> str:
    return dt.strftime("%d.%m.%Y %H:%M") if dt else ""


async def export_clients_xlsx(session: AsyncSession) -> bytes:
    clients = (
        await session.scalars(select(Client).options(selectinload(Client.cases)).order_by(Client.full_name))
    ).all()
    parties = (
        await session.scalars(select(Party).options(selectinload(Party.case)).order_by(Party.name))
    ).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Клієнти"
    ws.append(TEMPLATE_COLUMNS + ["Telegram", "Звернень", "Створено", "Джерело"])
    for c in clients:
        if not c.full_name:
            continue
        ws.append(
            [
                c.full_name,
                "юрособа" if c.is_company else "фізособа",
                c.phone,
                crypto.decrypt(c.tax_id_enc),
                "клієнт",
                "",
                c.note,
                f"@{c.telegram_username}" if c.telegram_username else "",
                len(c.cases),
                _fmt(c.created_at),
                {"bot": "бот", "import": "імпорт", "manual": "вручну"}.get(c.source, c.source),
            ]
        )
    for p in parties:
        related = p.related_client_name or (f"Звернення #{p.case_id}" if p.case_id else "")
        ws.append(
            [
                p.name,
                "юрособа" if p.is_company else "фізособа",
                "",
                "",
                p.role,
                related,
                p.note,
                "",
                "",
                _fmt(p.created_at),
                "звернення" if p.case_id else "імпорт",
            ]
        )
    _style_sheet(ws, [34, 14, 18, 16, 22, 30, 30, 18, 10, 18, 12])
    return _to_bytes(wb)


async def export_cases_xlsx(session: AsyncSession) -> bytes:
    cases = (
        await session.scalars(
            select(Case)
            .options(selectinload(Case.client), selectinload(Case.drafts))
            .order_by(Case.id.desc())
        )
    ).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Звернення"
    ws.append(
        [
            "№",
            "Створено",
            "Клієнт",
            "Телефон",
            "Назва",
            "Категорія",
            "Статус",
            "Терміново",
            "Складна",
            "Конфлікт",
            "Строк відповіді",
            "Консультацію надіслано",
        ]
    )
    for c in cases:
        sent = next((d.sent_at for d in c.drafts if d.sent_at), None)
        ws.append(
            [
                c.id,
                _fmt(c.created_at),
                c.client.full_name,
                c.client.phone,
                c.title,
                CATEGORIES.get(c.category, c.category),
                CaseStatus(c.status).label,
                "так" if c.is_urgent else "",
                "так" if c.is_complex else "",
                ConflictStatus(c.conflict_status).label,
                _fmt(c.response_due),
                _fmt(sent),
            ]
        )
    _style_sheet(ws, [6, 17, 30, 16, 40, 22, 22, 10, 9, 26, 17, 20])
    return _to_bytes(wb)
